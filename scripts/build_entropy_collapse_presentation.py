#!/usr/bin/env python3
"""Build the entropy collapse scaling presentation deck.

Expected environment:
    python-pptx
    matplotlib
    pandas
    numpy
    Pillow

This script generates:
    - slides/entropy-collapse-scaling-presentation.pptx
    - slides/entropy-collapse-scaling-presentation_outline.md
    - slides/entropy-collapse-scaling-presentation_manifest.md
    - slides/generated_assets/entropy-collapse-scaling/*.png
"""

from __future__ import annotations

import json
import math
import textwrap
from dataclasses import dataclass
import os
from pathlib import Path
import tempfile
from typing import Iterable

os.environ.setdefault("MPLBACKEND", "Agg")
os.environ.setdefault("MPLCONFIGDIR", str(Path(tempfile.gettempdir()) / "moltbook-mplconfig"))
os.environ.setdefault("XDG_CACHE_HOME", tempfile.gettempdir())

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
FINDINGS = ROOT / "findings" / "entropy-collapse-scaling"
STYLE_GUIDE = ROOT / "slides" / "fomat_details"
SLIDES_DIR = ROOT / "slides"
ASSET_DIR = SLIDES_DIR / "generated_assets" / "entropy-collapse-scaling"
PPTX_OUT = SLIDES_DIR / "entropy-collapse-scaling-presentation.pptx"
OUTLINE_OUT = SLIDES_DIR / "entropy-collapse-scaling-presentation_outline.md"
MANIFEST_OUT = SLIDES_DIR / "entropy-collapse-scaling-presentation_manifest.md"


TITLE_X = 0.26
TITLE_Y = 0.24
TITLE_W = 9.32
TITLE_H = 0.63

FONT_BODY = "Lato"
FONT_BOLD = "Lato Black"
FONT_FALLBACK = "Arial"

WHITE = "FFFFFF"
BLACK = "111111"
GRAY_900 = "212121"
GRAY_700 = "4B5563"
GRAY_600 = "6B7280"
GRAY_500 = "9CA3AF"
GRAY_300 = "D1D5DB"
GRAY_200 = "E5E7EB"
GRAY_100 = "F3F4F6"
BLUE = "2563EB"
BLUE_2 = "3B82F6"
GREEN = "059669"
RED = "DC2626"
ORANGE = "F97316"
YELLOW = "EAB308"
PINK = "E11D48"
NAVY = "0F172A"
PURPLE = "7C3AED"

CONDITION_COLORS = {
    "mag0": "6B7280",
    "mag1": "E11D48",
    "mag5": "F97316",
    "mag25": "EAB308",
    "dom-agi": "2563EB",
    "dom-tech": "059669",
}

CONDITION_LABELS = {
    "mag0": "Empty feed",
    "mag1": "1 conspiracy",
    "mag5": "5 conspiracies",
    "mag25": "25 conspiracies",
    "dom-agi": "25 AGI hype",
    "dom-tech": "25 tech humor",
}

SCALE_COLORS = {
    "n10": "94A3B8",
    "n20": "2563EB",
    "n30": "0F766E",
}

MODEL_CARD_ACCENTS = {
    "gpt-5": BLUE,
    "gemini-flash-lite": ORANGE,
    "glm-5": PURPLE,
    "kimi-k2.5": GREEN,
}


def hex_rgb(hex_code: str) -> RGBColor:
    hex_code = hex_code.replace("#", "")
    return RGBColor(int(hex_code[0:2], 16), int(hex_code[2:4], 16), int(hex_code[4:6], 16))


def lighten(hex_code: str, factor: float = 0.85) -> str:
    hex_code = hex_code.replace("#", "")
    r = int(hex_code[0:2], 16)
    g = int(hex_code[2:4], 16)
    b = int(hex_code[4:6], 16)
    r = int(r + (255 - r) * factor)
    g = int(g + (255 - g) * factor)
    b = int(b + (255 - b) * factor)
    return f"{r:02X}{g:02X}{b:02X}"


def short_phrase(text: str, max_len: int = 32) -> str:
    return text if len(text) <= max_len else text[: max_len - 1] + "…"


def ensure_dirs() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)


def setup_matplotlib() -> None:
    plt.style.use("default")
    plt.rcParams.update(
        {
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
            "font.size": 11,
            "axes.titlesize": 14,
            "axes.labelsize": 11,
            "xtick.labelsize": 9,
            "ytick.labelsize": 9,
            "font.family": "DejaVu Sans",
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.edgecolor": "#C8CDD5",
            "grid.color": "#E8ECF2",
            "grid.linewidth": 0.8,
        }
    )


@dataclass
class SlideMeta:
    number: int
    title: str
    key_message: str
    visual_spec: str
    notes: list[str]
    sources: list[str]


def load_json(model: str, rel: str) -> dict:
    return json.loads((FINDINGS / model / rel).read_text())


def image_size(path: Path) -> tuple[int, int]:
    with Image.open(path) as img:
        return img.size


def add_textbox(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    font_size: int = 16,
    bold: bool = False,
    color: str = BLACK,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    font_name: str | None = None,
    margin: float = 0.06,
    italic: bool = False,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin * 0.85)
    tf.margin_bottom = Inches(margin * 0.75)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    font = run.font
    font.name = font_name or (FONT_BOLD if bold else FONT_BODY)
    font.size = Pt(font_size)
    font.bold = bold
    font.italic = italic
    font.color.rgb = hex_rgb(color)


def add_paragraphs(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    paragraphs: Iterable[tuple[str, int, bool, str]],
    *,
    margin: float = 0.08,
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.clear()
    first = True
    for text, size, bold, color in paragraphs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = text
        run.font.name = FONT_BOLD if bold else FONT_BODY
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = hex_rgb(color)


def add_card(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    border: str = GRAY_300,
    fill: str = WHITE,
    radius: MSO_AUTO_SHAPE_TYPE = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
    line_width: float = 0.9,
) -> object:
    shape = slide.shapes.add_shape(radius, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = hex_rgb(fill)
    shape.line.color.rgb = hex_rgb(border)
    shape.line.width = Pt(line_width)
    return shape


def add_chip(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    color: str,
    *,
    text_color: str | None = None,
    font_size: int = 10,
    border: bool = False,
) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = hex_rgb(lighten(color, 0.82))
    if border:
        shape.line.color.rgb = hex_rgb(color)
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = FONT_BODY
    run.font.size = Pt(font_size)
    run.font.bold = True
    run.font.color.rgb = hex_rgb(text_color or color)


def add_title(slide, title: str, number: int, subtitle: str | None = None) -> None:
    add_textbox(
        slide,
        TITLE_X,
        TITLE_Y,
        TITLE_W,
        TITLE_H,
        title,
        font_size=24,
        bold=True,
        color=GRAY_900,
        font_name=FONT_BOLD,
    )
    if subtitle:
        add_textbox(
            slide,
            TITLE_X,
            TITLE_Y + 0.42,
            8.5,
            0.32,
            subtitle,
            font_size=10,
            color=GRAY_600,
            margin=0.0,
        )
    add_textbox(
        slide,
        9.28,
        5.18,
        0.45,
        0.22,
        f"‹{number}›",
        font_size=9,
        color=GRAY_500,
        align=PP_ALIGN.RIGHT,
        margin=0.0,
    )


def add_cover_title(slide, title: str, subtitle: str, number: int) -> None:
    add_textbox(
        slide,
        0.7,
        1.18,
        8.6,
        1.45,
        title,
        font_size=29,
        bold=True,
        color=GRAY_900,
        align=PP_ALIGN.CENTER,
        font_name=FONT_BOLD,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_textbox(
        slide,
        1.0,
        2.7,
        8.0,
        0.6,
        subtitle,
        font_size=16,
        color=GRAY_700,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(3.85), Inches(3.55), Inches(2.3), Inches(0.34)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = hex_rgb(lighten(BLUE, 0.88))
    shape.line.fill.background()
    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "Controlled Moltbook simulation"
    run.font.name = FONT_BODY
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = hex_rgb(BLUE)
    add_textbox(
        slide,
        9.28,
        5.18,
        0.45,
        0.22,
        f"‹{number}›",
        font_size=9,
        color=GRAY_500,
        align=PP_ALIGN.RIGHT,
        margin=0.0,
    )


def add_picture_fit(slide, path: Path, x: float, y: float, w: float, h: float) -> None:
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_post_card(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    agent: str,
    minute: float,
    title: str,
    phrase_label: str,
    accent: str,
) -> None:
    add_card(slide, x, y, w, h, border=GRAY_300, fill=WHITE, line_width=0.85)
    add_textbox(slide, x + 0.08, y + 0.06, w - 0.5, 0.18, f"u/{agent}", font_size=9, bold=True, color=GRAY_700)
    add_textbox(
        slide,
        x + w - 0.48,
        y + 0.06,
        0.35,
        0.18,
        f"{minute:04.1f}",
        font_size=8,
        color=GRAY_500,
        align=PP_ALIGN.RIGHT,
        margin=0.0,
    )
    add_textbox(
        slide,
        x + 0.08,
        y + 0.26,
        w - 0.18,
        0.63,
        title,
        font_size=15,
        bold=True,
        color=GRAY_900,
        font_name=FONT_BOLD,
    )
    add_chip(slide, x + 0.08, y + h - 0.28, w - 0.16, 0.18, phrase_label, accent, font_size=8, border=False)


def add_callout(slide, x: float, y: float, w: float, h: float, title: str, lines: list[str], accent: str = BLUE) -> None:
    add_card(slide, x, y, w, h, border=lighten(accent, 0.35), fill=lighten(accent, 0.92), line_width=0.9)
    add_textbox(slide, x + 0.08, y + 0.07, w - 0.16, 0.22, title, font_size=12, bold=True, color=accent)
    y_cursor = y + 0.33
    for line in lines:
        add_textbox(slide, x + 0.08, y_cursor, w - 0.18, 0.28, f"• {line}", font_size=10, color=GRAY_900)
        y_cursor += 0.25


def make_hero_scale_story() -> Path:
    out = ASSET_DIR / "hero_scale_story.png"
    df = pd.read_csv(FINDINGS / "gpt-5" / "diversity" / "diversity_metrics.csv")
    agg = (
        df.groupby(["scale", "bin_idx"], as_index=False)["distinct_5_cumulative"]
        .mean()
        .sort_values(["scale", "bin_idx"])
    )
    summary = load_json("gpt-5", "participation/participation_summary.json")
    per_scale = summary["concentration"]["per_scale"]

    bin_labels = ["0–15", "15–30", "30–45", "45–60"]

    fig, axes = plt.subplots(1, 2, figsize=(10.0, 3.2), dpi=220)
    ax = axes[0]
    for scale in ["n10", "n20", "n30"]:
        part = agg[agg["scale"] == scale]
        ax.plot(
            bin_labels,
            part["distinct_5_cumulative"],
            marker="o",
            linewidth=2.6,
            color="#" + SCALE_COLORS[scale],
            label=scale.replace("n", "") + " agents",
        )
        last_val = float(part["distinct_5_cumulative"].iloc[-1])
        ax.text(3.05, last_val, f"{last_val:.2f}", fontsize=9, color="#" + SCALE_COLORS[scale], va="center")
    ax.set_ylim(0.74, 0.98)
    ax.set_title("GPT-5 average cumulative 5-gram diversity")
    ax.set_ylabel("Distinct 5-grams")
    ax.grid(axis="y")
    ax.legend(frameon=False, loc="lower left")

    ax = axes[1]
    scales = ["n10", "n20", "n30"]
    adoption = [per_scale[s]["mean_adoption_rate"] for s in scales]
    adopters = [per_scale[s]["mean_adopters"] for s in scales]
    bars = ax.bar(
        ["10", "20", "30"],
        adoption,
        color=["#" + SCALE_COLORS[s] for s in scales],
        width=0.6,
    )
    for bar, share, adopters_count in zip(bars, adoption, adopters):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            share + 0.02,
            f"{share:.2f}\n({adopters_count:.1f} agents)",
            ha="center",
            va="bottom",
            fontsize=9,
            color="#111111",
        )
    ax.set_ylim(0, 0.66)
    ax.set_title("Top phrase still reaches a large share of the run")
    ax.set_ylabel("Mean adoption rate")
    ax.set_xlabel("Agent count")
    ax.grid(axis="y")

    fig.tight_layout()
    fig.savefig(out, bbox_inches="tight")
    plt.close(fig)
    return out


def make_domtech_dynamics() -> Path:
    out = ASSET_DIR / "domtech_dynamics.png"
    df = pd.read_csv(FINDINGS / "gpt-5" / "diffusion" / "first_usage_timeline.csv")
    df = df[(df["condition"] == "dom-tech") & (df["phrase_rank"] == 1)].copy()
    fig, ax = plt.subplots(figsize=(9.2, 3.1), dpi=220)
    phrase_labels = []
    for scale in ["n10", "n20", "n30"]:
        part = df[df["scale"] == scale].sort_values("first_minute")
        part["cum"] = np.arange(1, len(part) + 1) / part["total_agents_in_run"].iloc[0]
        x = [0.0] + part["first_minute"].tolist() + [60.0]
        y = [0.0] + part["cum"].tolist() + [part["cum"].iloc[-1]]
        ax.step(x, y, where="post", linewidth=2.6, color="#" + SCALE_COLORS[scale], label=scale.replace("n", "") + " agents")
        phrase = short_phrase(part["phrase"].iloc[0], 28)
        phrase_labels.append(f"{scale.replace('n','')}a: {phrase}")
        ax.text(60.3, y[-1], f"{int(round(y[-1]*100))}%", fontsize=9, color="#" + SCALE_COLORS[scale], va="center")
    ax.set_xlim(0, 60)
    ax.set_ylim(0, 0.82)
    ax.set_xlabel("Minutes")
    ax.set_ylabel("Share of agents using that run's #1 phrase")
    ax.grid(axis="y")
    ax.legend(frameon=False, loc="upper left")
    fig.text(0.13, -0.02, " | ".join(phrase_labels), fontsize=9, color="#4B5563")
    fig.tight_layout()
    fig.savefig(out, bbox_inches="tight")
    plt.close(fig)
    return out


def make_collective_effect() -> Path:
    out = ASSET_DIR / "collective_effect.png"
    summary = load_json("gpt-5", "participation/participation_summary.json")
    per_scale = summary["concentration"]["per_scale"]
    overlap = summary["phrase_overlap"]["per_scale"]
    scales = ["n10", "n20", "n30"]

    fig, axes = plt.subplots(1, 3, figsize=(10.5, 3.0), dpi=220)
    x = np.arange(3)

    adopters = [per_scale[s]["mean_adopters"] for s in scales]
    axes[0].bar(x, adopters, color=["#" + SCALE_COLORS[s] for s in scales], width=0.6)
    axes[0].set_title("More agents adopt the dominant phrase")
    axes[0].set_xticks(x, ["10", "20", "30"])
    axes[0].set_ylabel("Mean adopters")
    axes[0].grid(axis="y")
    for idx, val in enumerate(adopters):
        axes[0].text(idx, val + 0.4, f"{val:.1f}", ha="center", fontsize=9)

    top1 = [per_scale[s]["mean_top1_share"] for s in scales]
    axes[1].bar(x, top1, color=["#" + SCALE_COLORS[s] for s in scales], width=0.6)
    axes[1].set_title("Single-agent dominance falls")
    axes[1].set_xticks(x, ["10", "20", "30"])
    axes[1].set_ylim(0, 0.65)
    axes[1].set_ylabel("Mean top-1 agent share")
    axes[1].grid(axis="y")
    for idx, val in enumerate(top1):
        axes[1].text(idx, val + 0.03, f"{val:.2f}", ha="center", fontsize=9)

    jaccard = [overlap[s]["mean_jaccard"] for s in scales]
    axes[2].bar(x, jaccard, color=["#" + SCALE_COLORS[s] for s in scales], width=0.6)
    axes[2].set_title("Top phrases share many of the same adopters")
    axes[2].set_xticks(x, ["10", "20", "30"])
    axes[2].set_ylim(0, 0.9)
    axes[2].set_ylabel("Mean adopter overlap (Jaccard)")
    axes[2].grid(axis="y")
    for idx, val in enumerate(jaccard):
        axes[2].text(idx, val + 0.03, f"{val:.2f}", ha="center", fontsize=9)

    fig.tight_layout()
    fig.savefig(out, bbox_inches="tight")
    plt.close(fig)
    return out


def make_semantic_concentration() -> Path:
    out = ASSET_DIR / "semantic_concentration.png"
    df = pd.read_csv(FINDINGS / "gpt-5" / "embedding_bridge" / "topic_anchor_run_summary.csv")
    chosen = df[df["condition"].isin(["dom-agi", "dom-tech", "mag25"])].copy()
    chosen["label"] = chosen.apply(
        lambda row: f"{row['scale'].replace('n','')}a {CONDITION_LABELS[row['condition']]} [{row['family_dominant_topic']}]",
        axis=1,
    )
    chosen = chosen.sort_values(["condition", "scale"])
    y = np.arange(len(chosen))

    fig, ax = plt.subplots(figsize=(8.8, 4.0), dpi=220)
    ax.barh(y + 0.16, chosen["all_dominant_share"], height=0.28, color="#D6DBE5", label="All posts")
    colors = ["#" + CONDITION_COLORS[c] for c in chosen["condition"]]
    ax.barh(y - 0.16, chosen["family_dominant_share"], height=0.28, color=colors, label="Dominant phrase-family posts")
    for idx, row in enumerate(chosen.itertuples(index=False)):
        ax.text(row.all_dominant_share + 0.01, idx + 0.16, f"{row.all_dominant_share:.2f}", fontsize=8, va="center")
        ax.text(row.family_dominant_share + 0.01, idx - 0.16, f"{row.family_dominant_share:.2f}", fontsize=8, va="center")
    ax.set_yticks(y, chosen["label"])
    ax.set_xlim(0, 1.05)
    ax.set_xlabel("Share of posts near the run's dominant topic anchor")
    ax.grid(axis="x")
    ax.legend(frameon=False, loc="lower left", bbox_to_anchor=(0.0, 1.01), ncol=2)
    fig.tight_layout()
    fig.savefig(out, bbox_inches="tight")
    plt.close(fig)
    return out


def get_total_runs() -> dict[str, int]:
    total = {}
    for model in ["gpt-5", "gemini-flash-lite", "glm-5", "kimi-k2.5"]:
        total[model] = len(load_json(model, "diffusion/diffusion_summary.json")["per_run"])
    return total


def get_opening_posts() -> tuple[list[dict], dict]:
    examples = load_json("gpt-5", "participation/example_posts.json")
    key = "n30/dom-tech/receipt why options owner link"
    posts = examples[key][:4]
    summary = load_json("gpt-5", "diffusion/diffusion_summary.json")["per_run"]["n30/dom-tech"]["#1"]
    return posts, summary


def get_origin_examples() -> list[dict]:
    diffusion = load_json("gpt-5", "diffusion/diffusion_summary.json")["per_run"]
    runs = ["n10/mag0", "n10/dom-agi", "n20/mag5", "n20/dom-tech", "n30/mag25", "n30/dom-tech"]
    items = []
    for run in runs:
        scale, condition = run.split("/")
        info = diffusion[run]["#1"]
        items.append(
            {
                "run": run,
                "scale": scale,
                "condition": condition,
                "label": f"{scale.replace('n','')}a {CONDITION_LABELS[condition]}",
                "phrase": info["phrase"],
                "rate": info["rate"],
            }
        )
    return items


def get_cross_condition_examples() -> list[dict]:
    labels_30 = {
        "mag0": "Small-bet accountability ritual",
        "mag1": "Checklist epistemics for posting",
        "mag5": "Calm epistemic checklist for hot threads",
        "mag25": "Receipt-first micro-science template",
        "dom-agi": "Receipts-and-rollback safety checklist",
        "dom-tech": "Calm-by-design legibility micro-pilot",
    }
    diffusion = load_json("gpt-5", "diffusion/diffusion_summary.json")["per_run"]
    topic = pd.read_csv(FINDINGS / "gpt-5" / "embedding_bridge" / "topic_anchor_run_summary.csv")
    topic = topic[topic["scale"] == "n30"].set_index("condition")
    items = []
    for condition in ["mag0", "mag1", "mag5", "mag25", "dom-agi", "dom-tech"]:
        info = diffusion[f"n30/{condition}"]["#1"]
        topic_row = topic.loc[condition]
        items.append(
            {
                "condition": condition,
                "condition_label": CONDITION_LABELS[condition],
                "phrase": info["phrase"],
                "rate": info["rate"],
                "anchor": topic_row["all_dominant_topic"],
                "label": labels_30[condition],
            }
        )
    return items


def get_model_cards() -> list[dict]:
    cards = []
    for model in ["gpt-5", "gemini-flash-lite", "glm-5", "kimi-k2.5"]:
        participation = load_json(model, "participation/participation_summary.json")
        provenance = load_json(model, "provenance/provenance_summary.json")
        scales = list(participation["concentration"]["per_scale"].keys())
        n10 = participation["concentration"]["per_scale"]["n10"]
        cards.append(
            {
                "model": model,
                "scales": scales,
                "mean_adoption_n10": n10["mean_adoption_rate"],
                "mean_adopters_n10": n10["mean_adopters"],
                "fivegram_overlap": provenance["per_ngram"]["5"]["mean_cross_run_overlap"],
            }
        )
    return cards


def register(
    metas: list[SlideMeta],
    number: int,
    title: str,
    key_message: str,
    visual_spec: str,
    notes: list[str],
    sources: list[str],
) -> None:
    metas.append(SlideMeta(number, title, key_message, visual_spec, notes, sources))


def build_deck() -> list[SlideMeta]:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    blank = prs.slide_layouts[6]

    setup_matplotlib()
    ensure_dirs()

    hero_chart = make_hero_scale_story()
    domtech_chart = make_domtech_dynamics()
    collective_chart = make_collective_effect()
    semantic_chart = make_semantic_concentration()

    metas: list[SlideMeta] = []
    opening_posts, opening_summary = get_opening_posts()
    origin_examples = get_origin_examples()
    cross_condition_examples = get_cross_condition_examples()
    model_cards = get_model_cards()
    total_runs = get_total_runs()

    # Slide 1
    slide = prs.slides.add_slide(blank)
    add_cover_title(
        slide,
        "What Happens When AI Agents Get a Social Feed?",
        "Moltbook gave the world the viral demo. We ran the controlled version.",
        1,
    )
    register(
        metas,
        1,
        "What Happens When AI Agents Get a Social Feed?",
        "Use Moltbook as the public motivation, then pivot immediately to the controlled experiment.",
        "Centered title slide with one short subtitle and a single control-simulation chip.",
        [
            "Moltbook made agent social media feel real to a broad audience.",
            "But the research question needs cleaner measurement than the live public platform could provide.",
            "This talk asks what agent discourse actually does under controlled social conditions.",
        ],
        [str(STYLE_GUIDE / "README.md")],
    )

    # Slide 2
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Moltbook was exciting. It was also a bad measurement environment.", 2)
    add_paragraphs(
        slide,
        0.48,
        1.1,
        4.25,
        2.3,
        [
            ("Moltbook became the public proof-of-concept for AI-agent social media.", 17, True, GRAY_900),
            (
                "But once human prompting, impersonation, and manipulation entered the loop, it stopped being a clean place to measure what autonomous agent interaction was actually doing.",
                13,
                False,
                GRAY_700,
            ),
            ("So we kept the social setting and rebuilt the experiment under controlled conditions.", 13, False, GRAY_700),
        ],
    )
    for idx, (label, accent, detail) in enumerate(
        [
            ("Viral public demo", BLUE, "The internet already cared."),
            ("Too much ambiguity", ORANGE, "Human influence blurred the signal."),
            ("Controlled rebuild", GREEN, "Same question, cleaner measurement."),
        ]
    ):
        yy = 1.2 + idx * 1.02
        add_card(slide, 5.45, yy, 3.82, 0.8, border=lighten(accent, 0.4), fill=lighten(accent, 0.93))
        add_textbox(slide, 5.65, yy + 0.11, 3.2, 0.23, label, font_size=14, bold=True, color=accent)
        add_textbox(slide, 5.65, yy + 0.38, 3.1, 0.22, detail, font_size=10, color=GRAY_700)
    register(
        metas,
        2,
        "Moltbook was exciting. It was also a bad measurement environment.",
        "The public platform motivates the question, but the evidence must come from the controlled simulation.",
        "Left text block plus three compact cards: viral demo, ambiguity, controlled rebuild.",
        [
            "Use the public story to explain why this matters, not as the evidence base.",
            "Notes can cite the January 31, 2026 Fortune piece, the February 12, 2026 Euronews reporting, and the March 10, 2026 TechCrunch coverage.",
            "Keep the on-screen wording general: human prompting, impersonation, and manipulation made it noisy.",
        ],
        [
            "https://fortune.com/2026/01/31/ai-agent-moltbot-clawdbot-openclaw-data-privacy-security-nightmare-moltbook-social-network//",
            "https://www.euronews.com/next/2026/02/12/ai-or-human-researchers-question-whos-posting-on-ai-bot-social-media-site-moltbook",
            "https://techcrunch.com/2026/03/10/meta-acquired-moltbook-the-ai-agent-social-network-that-went-viral-because-of-fake-posts/",
        ],
    )

    # Slide 3
    slide = prs.slides.add_slide(blank)
    add_title(
        slide,
        "Different agents start posting like the same kind of account.",
        3,
        "Representative GPT-5 posts from one 30-agent tech-humor run",
    )
    positions = [(0.42, 1.02), (5.04, 1.02), (0.42, 3.05), (5.04, 3.05)]
    phrase_chip = "receipt / why / options / owner / link"
    for post, (x, y) in zip(opening_posts, positions):
        add_post_card(
            slide,
            x,
            y,
            4.28,
            1.78,
            agent=post["agent"],
            minute=post["minute"],
            title=post["title"],
            phrase_label=phrase_chip,
            accent=CONDITION_COLORS["dom-tech"],
        )
    add_callout(
        slide,
        7.38,
        4.73,
        2.12,
        0.55,
        "Same run",
        [f"{opening_summary['adopters']}/{opening_summary['total_agents']} agents used the dominant family"],
        accent=CONDITION_COLORS["dom-tech"],
    )
    register(
        metas,
        3,
        "Different agents start posting like the same kind of account.",
        "Establish the symptom in plain human terms before any metric-heavy slide appears.",
        "2x2 grid of Moltbook-style post cards from one run, plus one small adoption callout.",
        [
            "The point is not that the titles are identical; it is that they feel like one shared account voice or house style.",
            "All four cards are from different agents inside the same run.",
            f"The dominant phrase family in this run reached {opening_summary['adopters']}/{opening_summary['total_agents']} agents and {opening_summary['count']} total uses.",
        ],
        [
            str(FINDINGS / "gpt-5" / "participation" / "example_posts.json"),
            str(FINDINGS / "gpt-5" / "diffusion" / "diffusion_summary.json"),
        ],
    )

    # Slide 4
    slide = prs.slides.add_slide(blank)
    add_title(slide, "So we rebuilt the setting under controlled conditions.", 4)
    add_card(slide, 0.55, 1.25, 8.85, 2.15, border=GRAY_200, fill=WHITE)
    stages = [
        ("Moltbook-like feed", "Same social affordance:\nposts, replies, voting", BLUE),
        ("Agent population", "Distinct archetypes,\nshared feed exposure", ORANGE),
        ("Controlled logging", "Every post, phrase,\nand adoption event traced", GREEN),
    ]
    x_positions = [0.88, 3.62, 6.36]
    for (title, detail, accent), x in zip(stages, x_positions):
        add_card(slide, x, 1.78, 2.0, 1.05, border=lighten(accent, 0.35), fill=lighten(accent, 0.93))
        add_textbox(slide, x + 0.08, 1.9, 1.8, 0.2, title, font_size=13, bold=True, color=accent)
        add_textbox(slide, x + 0.08, 2.18, 1.8, 0.42, detail, font_size=9, color=GRAY_700)
    for x1, x2 in zip([2.9, 5.64], [3.62, 6.36]):
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(2.3), Inches(x2), Inches(2.3)
        )
        line.line.color.rgb = hex_rgb(GRAY_300)
        line.line.width = Pt(1.0)
    add_paragraphs(
        slide,
        0.62,
        3.75,
        8.4,
        1.0,
        [
            ("We kept the social setting and the seed worlds.", 15, True, GRAY_900),
            (
                "We removed the public noise so we could measure whether the agents become richer or more predictable when they interact at scale.",
                12,
                False,
                GRAY_700,
            ),
        ],
    )
    register(
        metas,
        4,
        "So we rebuilt the setting under controlled conditions.",
        "Bridge from the public Moltbook story to the actual evidence base for the talk.",
        "One wide pipeline diagram plus a two-line explanatory block.",
        [
            "This is the transition slide. Its job is only to say why the controlled simulation exists.",
            "Do not overload it with metrics.",
        ],
        [],
    )

    # Slide 5
    slide = prs.slides.add_slide(blank)
    add_title(slide, "What we wanted to know.", 5)
    add_card(slide, 0.62, 1.15, 8.7, 1.12, border=lighten(BLUE, 0.5), fill=lighten(BLUE, 0.94))
    add_textbox(
        slide,
        0.86,
        1.38,
        8.15,
        0.58,
        "When LLM agents interact freely on a Reddit-like social platform, do they develop richer social behavior or narrow into repeatable patterns?",
        font_size=20,
        bold=True,
        color=GRAY_900,
        font_name=FONT_BOLD,
        valign=MSO_ANCHOR.MIDDLE,
    )
    extras = [
        ("Does scale help?", "Do 20 or 30 agents make the discourse more open-ended?"),
        ("Is it collective?", "Or is the effect just one or two loud accounts?"),
    ]
    for idx, (title, detail) in enumerate(extras):
        x = 0.82 + idx * 4.3
        add_card(slide, x, 2.85, 3.95, 1.45, border=GRAY_200, fill=WHITE)
        add_textbox(slide, x + 0.1, 3.02, 3.4, 0.22, title, font_size=16, bold=True, color=GRAY_900)
        add_textbox(slide, x + 0.1, 3.35, 3.4, 0.52, detail, font_size=12, color=GRAY_700)
    register(
        metas,
        5,
        "What we wanted to know.",
        "Lock the main question and the two supporting questions before moving into setup and results.",
        "One large question card and two smaller supporting-question cards.",
        [
            "This slide should sound plainspoken, not paper-like.",
            "The three question structure becomes the deck spine.",
        ],
        [],
    )

    # Slide 6
    slide = prs.slides.add_slide(blank)
    add_title(slide, "The controlled study stayed simple on purpose.", 6)
    add_card(slide, 0.55, 1.05, 8.95, 3.05, border=GRAY_200, fill=WHITE)
    setup_cards = [
        ("6 feed conditions", "Empty feed, 3 conspiracy doses, AGI hype, tech humor", GRAY_700),
        ("10 / 20 / 30 agents", "Two full scaling models plus two partial supporting replications", BLUE),
        ("4 models", "GPT-5, Gemini Flash Lite, GLM-5, Kimi-K2.5", ORANGE),
        ("1-hour runs", "Every post, phrase, and adoption event logged", GREEN),
    ]
    coords = [(0.82, 1.35), (5.05, 1.35), (0.82, 2.38), (5.05, 2.38)]
    for (title, detail, accent), (x, y) in zip(setup_cards, coords):
        add_card(slide, x, y, 3.55, 0.82, border=lighten(accent, 0.4), fill=lighten(accent, 0.94))
        add_textbox(slide, x + 0.1, y + 0.1, 3.05, 0.2, title, font_size=13, bold=True, color=accent)
        add_textbox(slide, x + 0.1, y + 0.37, 3.1, 0.25, detail, font_size=9, color=GRAY_700)
    add_chip(slide, 0.8, 4.35, 1.0, 0.2, "mag0", CONDITION_COLORS["mag0"], font_size=8)
    add_chip(slide, 1.95, 4.35, 1.0, 0.2, "mag1", CONDITION_COLORS["mag1"], font_size=8)
    add_chip(slide, 3.1, 4.35, 1.0, 0.2, "mag5", CONDITION_COLORS["mag5"], font_size=8)
    add_chip(slide, 4.25, 4.35, 1.0, 0.2, "mag25", CONDITION_COLORS["mag25"], font_size=8)
    add_chip(slide, 5.42, 4.35, 1.2, 0.2, "dom-agi", CONDITION_COLORS["dom-agi"], font_size=8)
    add_chip(slide, 6.8, 4.35, 1.3, 0.2, "dom-tech", CONDITION_COLORS["dom-tech"], font_size=8)
    add_textbox(
        slide,
        0.82,
        4.66,
        7.85,
        0.22,
        f"Runs in this folder: GPT-5 {total_runs['gpt-5']}, Gemini Flash Lite {total_runs['gemini-flash-lite']}, GLM-5 {total_runs['glm-5']}, Kimi-K2.5 {total_runs['kimi-k2.5']}.",
        font_size=10,
        color=GRAY_600,
        margin=0.0,
    )
    register(
        metas,
        6,
        "The controlled study stayed simple on purpose.",
        "Show the experimental degrees of freedom without slowing the deck down.",
        "Four setup cards, one condition strip, one run-count line.",
        [
            "Keep this brisk because the audience already knows the platform context.",
            "The slide is here so later results never feel under-defined.",
        ],
        [
            str(FINDINGS / "gpt-5" / "diffusion" / "diffusion_summary.json"),
            str(FINDINGS / "gemini-flash-lite" / "diffusion" / "diffusion_summary.json"),
            str(FINDINGS / "glm-5" / "diffusion" / "diffusion_summary.json"),
            str(FINDINGS / "kimi-k2.5" / "diffusion" / "diffusion_summary.json"),
        ],
    )

    # Slide 7
    slide = prs.slides.add_slide(blank)
    add_title(slide, "We measured whether the feed became narrower, stickier, and more shared.", 7)
    blocks = [
        ("Diversity loss", "Do the posts use fewer distinct 5-grams over time?", BLUE),
        ("Phrase diffusion", "Does one local phrase family spread through the run?", ORANGE),
        ("Collective adoption", "Is the dominant pattern shared across many agents?", GREEN),
    ]
    xs = [0.6, 3.43, 6.26]
    for (title, detail, accent), x in zip(blocks, xs):
        add_card(slide, x, 1.5, 2.62, 1.95, border=lighten(accent, 0.4), fill=lighten(accent, 0.94))
        add_textbox(slide, x + 0.12, 1.72, 2.2, 0.24, title, font_size=16, bold=True, color=accent)
        add_textbox(slide, x + 0.12, 2.08, 2.26, 0.92, detail, font_size=12, color=GRAY_700)
    add_card(slide, 1.2, 4.25, 7.6, 0.55, border=GRAY_200, fill=GRAY_100)
    add_textbox(
        slide,
        1.36,
        4.39,
        7.18,
        0.2,
        "The point is not just that language repeats. The point is that repetition becomes socially shared.",
        font_size=12,
        color=GRAY_700,
        align=PP_ALIGN.CENTER,
        margin=0.0,
    )
    register(
        metas,
        7,
        "We measured whether the feed became narrower, stickier, and more shared.",
        "Define the deck’s metric language in plain English before showing any figure.",
        "Three metric cards plus one bottom-line strip.",
        [
            "Use human language here, not method jargon.",
            "This slide should make the later charts self-explanatory.",
        ],
        [],
    )

    # Slide 8
    slide = prs.slides.add_slide(blank)
    add_title(slide, "More agents do not automatically make the discourse richer.", 8)
    add_picture_fit(slide, hero_chart, 0.55, 1.08, 6.45, 3.6)
    add_callout(
        slide,
        7.18,
        1.38,
        2.1,
        1.65,
        "GPT-5 scale signal",
        [
            "Late-window cumulative 5-gram diversity falls from 0.87 to 0.78.",
            "The top phrase still reaches 43% to 54% of agents on average.",
            "More agents do not create automatic takeoff into richer discourse.",
        ],
        accent=BLUE,
    )
    add_textbox(
        slide,
        7.22,
        3.35,
        2.0,
        0.75,
        "Gemini Flash Lite behaves differently on the surface, but larger populations still do not reliably open the space up.",
        font_size=10,
        color=GRAY_700,
    )
    register(
        metas,
        8,
        "More agents do not automatically make the discourse richer.",
        "Answer the scale question cleanly and early, using the clearest full-scaling result.",
        "One rebuilt two-panel figure plus a compact metric callout.",
        [
            "This is the first major answer slide.",
            "Lead with GPT-5 because the scale trend is easiest to read there.",
            "Mention Gemini as supporting contrast, not as the primary narrative driver here.",
        ],
        [
            str(FINDINGS / "gpt-5" / "diversity" / "diversity_metrics.csv"),
            str(FINDINGS / "gpt-5" / "participation" / "participation_summary.json"),
            str(FINDINGS / "gemini-flash-lite" / "diversity" / "diversity_metrics.csv"),
        ],
    )

    # Slide 9
    slide = prs.slides.add_slide(blank)
    add_title(slide, "A shared posting style emerges inside a run.", 9)
    add_picture_fit(slide, domtech_chart, 0.55, 1.05, 6.9, 3.55)
    add_callout(
        slide,
        7.55,
        1.3,
        1.72,
        1.75,
        "Running example",
        [
            "GPT-5 / 25 tech humor",
            "30 agents: 21/30 adopters",
            "383 total uses",
            "First appearance at minute 2.4",
        ],
        accent=CONDITION_COLORS["dom-tech"],
    )
    add_textbox(
        slide,
        7.55,
        3.35,
        1.72,
        0.72,
        "This is not one phrase echoing because one agent got loud. It keeps getting picked up by new agents.",
        font_size=10,
        color=GRAY_700,
    )
    register(
        metas,
        9,
        "A shared posting style emerges inside a run.",
        "Show the social spread of one local phrase family over time using a single clean running example.",
        "One rebuilt adoption-curve figure and one narrow callout card.",
        [
            "The narrative shift here is from static outcome to temporal spread.",
            "The chosen condition is deliberately memorable and visually strong.",
        ],
        [
            str(FINDINGS / "gpt-5" / "diffusion" / "first_usage_timeline.csv"),
            str(FINDINGS / "gpt-5" / "diffusion" / "diffusion_summary.json"),
        ],
    )

    # Slide 10
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Each run develops its own local phrase family.", 10)
    add_textbox(
        slide,
        0.58,
        0.74,
        4.8,
        0.22,
        "Representative GPT-5 top phrase families across runs",
        font_size=10,
        color=GRAY_600,
        margin=0.0,
    )
    positions = [(0.55, 1.02), (3.36, 1.02), (6.17, 1.02), (0.55, 3.02), (3.36, 3.02), (6.17, 3.02)]
    for item, (x, y) in zip(origin_examples, positions):
        accent = CONDITION_COLORS[item["condition"]]
        add_card(slide, x, y, 2.55, 1.62, border=lighten(accent, 0.45), fill=WHITE)
        add_chip(slide, x + 0.09, y + 0.1, 1.25, 0.18, item["label"], accent, font_size=8)
        add_textbox(
            slide,
            x + 0.11,
            y + 0.42,
            2.2,
            0.52,
            item["phrase"],
            font_size=13,
            bold=True,
            color=GRAY_900,
            font_name=FONT_BOLD,
        )
        add_textbox(
            slide,
            x + 0.11,
            y + 1.2,
            2.15,
            0.2,
            f"Top phrase adoption: {int(round(item['rate'] * 100))}%",
            font_size=9,
            color=GRAY_600,
            margin=0.0,
        )
    add_card(slide, 1.1, 4.88, 7.8, 0.28, border=GRAY_200, fill=GRAY_100)
    add_textbox(
        slide,
        1.22,
        4.93,
        7.55,
        0.16,
        "Same phenomenon, different local attractors: convergence is social within a run, not a single universal phrase across all runs.",
        font_size=10,
        color=GRAY_700,
        align=PP_ALIGN.CENTER,
        margin=0.0,
    )
    register(
        metas,
        10,
        "Each run develops its own local phrase family.",
        "Show that convergence is local to the run rather than a single repeated seed phrase everywhere.",
        "Six phrase-family cards spanning different runs and conditions, plus one bottom explanatory strip.",
        [
            "This is the clean answer to the seed-copying concern.",
            "Use the phrase cards as memorable artifacts, not as text the audience must read word-for-word.",
        ],
        [
            str(FINDINGS / "gpt-5" / "diffusion" / "diffusion_summary.json"),
            str(FINDINGS / "gpt-5" / "provenance" / "provenance_summary.json"),
        ],
    )

    # Slide 11
    slide = prs.slides.add_slide(blank)
    add_title(slide, "This is not one loud account.", 11)
    add_picture_fit(slide, collective_chart, 0.48, 1.1, 6.8, 3.45)
    add_callout(
        slide,
        7.38,
        1.32,
        1.92,
        1.8,
        "Population effect",
        [
            "Mean adopters rise from 4.3 to 16.3.",
            "Mean top-1 agent share falls from 0.51 to 0.22.",
            "Top phrases increasingly share the same adopters.",
        ],
        accent=GREEN,
    )
    add_textbox(
        slide,
        7.4,
        3.45,
        1.8,
        0.62,
        "As scale grows, the dominant pattern becomes more collective even while single-agent dominance weakens.",
        font_size=10,
        color=GRAY_700,
    )
    register(
        metas,
        11,
        "This is not one loud account.",
        "Answer the second major supporting question by showing the shift from individual dominance to shared adoption.",
        "One rebuilt three-panel collective-effect figure plus a compact metric callout.",
        [
            "This is the strongest anti-'few spammer accounts' slide in the deck.",
            "Keep the explanation causal and plainspoken.",
        ],
        [
            str(FINDINGS / "gpt-5" / "participation" / "participation_summary.json"),
            str(FINDINGS / "gpt-5" / "participation" / "concentration.csv"),
        ],
    )

    # Slide 12
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Different feeds change the vocabulary more than the social form.", 12)
    card_w = 1.42
    x_start = 0.52
    for idx, item in enumerate(cross_condition_examples):
        x = x_start + idx * 1.52
        accent = CONDITION_COLORS[item["condition"]]
        add_card(slide, x, 1.18, card_w, 2.95, border=lighten(accent, 0.45), fill=WHITE)
        add_chip(slide, x + 0.08, 1.28, 1.0, 0.18, item["condition_label"], accent, font_size=7)
        add_textbox(
            slide,
            x + 0.08,
            1.64,
            card_w - 0.16,
            0.48,
            item["phrase"],
            font_size=11,
            bold=True,
            color=GRAY_900,
            font_name=FONT_BOLD,
        )
        add_textbox(
            slide,
            x + 0.08,
            2.46,
            card_w - 0.16,
            0.44,
            item["label"],
            font_size=10,
            color=GRAY_700,
        )
        add_textbox(
            slide,
            x + 0.08,
            3.45,
            card_w - 0.16,
            0.2,
            f"Anchor: {item['anchor']} • {int(round(item['rate']*100))}% adoption",
            font_size=7,
            color=GRAY_600,
            margin=0.0,
        )
    add_card(slide, 0.86, 4.5, 8.1, 0.4, border=GRAY_200, fill=GRAY_100)
    add_textbox(
        slide,
        1.0,
        4.58,
        7.8,
        0.2,
        "The nouns change. The dominant social format stays procedural: checklists, receipts, loops, off-ramps, review dates.",
        font_size=11,
        color=GRAY_700,
        align=PP_ALIGN.CENTER,
        margin=0.0,
    )
    register(
        metas,
        12,
        "Different feeds change the vocabulary more than the social form.",
        "Show that different seed worlds steer the surface language, but many runs still narrow into procedural social formats.",
        "Six slim condition cards using representative n30 GPT-5 phrase families and concise human labels.",
        [
            "This slide is about form versus topic, not about claiming the conditions are identical.",
            "The summary labels come from the repo’s own cross-condition interpretation notes.",
        ],
        [
            str(FINDINGS / "gpt-5" / "diffusion" / "diffusion_summary.json"),
            str(ROOT / "findings" / "moltbook-entropy-collapse-scaling" / "summary.md"),
            str(FINDINGS / "gpt-5" / "embedding_bridge" / "topic_anchor_run_summary.csv"),
        ],
    )

    # Slide 13
    slide = prs.slides.add_slide(blank)
    add_title(slide, "The pattern is visible across models, not just one model family.", 13)
    positions = [(0.62, 1.08), (5.02, 1.08), (0.62, 3.15), (5.02, 3.15)]
    display_names = {
        "gpt-5": "GPT-5",
        "gemini-flash-lite": "Gemini Flash Lite",
        "glm-5": "GLM-5",
        "kimi-k2.5": "Kimi-K2.5",
    }
    for card, (x, y) in zip(model_cards, positions):
        accent = MODEL_CARD_ACCENTS[card["model"]]
        add_card(slide, x, y, 3.95, 1.48, border=lighten(accent, 0.42), fill=WHITE)
        add_textbox(slide, x + 0.1, y + 0.1, 2.4, 0.2, display_names[card["model"]], font_size=16, bold=True, color=accent)
        scale_text = " / ".join(s.replace("n", "") for s in card["scales"])
        add_textbox(slide, x + 0.1, y + 0.4, 1.4, 0.18, f"Scales here: {scale_text}", font_size=9, color=GRAY_700, margin=0.0)
        add_textbox(
            slide,
            x + 0.1,
            y + 0.7,
            3.45,
            0.18,
            f"Mean top-phrase adoption at n10: {card['mean_adoption_n10']:.2f}",
            font_size=9,
            color=GRAY_700,
            margin=0.0,
        )
        add_textbox(
            slide,
            x + 0.1,
            y + 0.95,
            3.55,
            0.18,
            f"5-gram cross-run overlap: {card['fivegram_overlap']:.2f}",
            font_size=9,
            color=GRAY_700,
            margin=0.0,
        )
        note = "Full 10/20/30 scaling in this deck" if len(card["scales"]) == 3 else "n10 in this dataset"
        add_chip(slide, x + 2.55, y + 0.1, 1.2, 0.18, note, accent, font_size=7)
    add_card(slide, 1.15, 4.78, 7.5, 0.24, border=GRAY_200, fill=GRAY_100)
    add_textbox(
        slide,
        1.28,
        4.82,
        7.2,
        0.14,
        "Broad claim: all four models show local phrase attractors. Strongest scale evidence here comes from GPT-5 and Gemini Flash Lite.",
        font_size=9,
        color=GRAY_700,
        align=PP_ALIGN.CENTER,
        margin=0.0,
    )
    register(
        metas,
        13,
        "The pattern is visible across models, not just one model family.",
        "Support the broad multi-model takeaway without overstating what is fully scaled in this folder.",
        "Four model cards with scale coverage, n10 adoption, and 5-gram overlap, plus one small footer note.",
        [
            "The key claim is modest and defensible: local phrase attractors appear across all four models.",
            "The stronger 10/20/30 scale claim should remain tied to GPT-5 and Gemini on this slide.",
        ],
        [
            str(FINDINGS / "gpt-5" / "participation" / "participation_summary.json"),
            str(FINDINGS / "gpt-5" / "provenance" / "provenance_summary.json"),
            str(FINDINGS / "gemini-flash-lite" / "participation" / "participation_summary.json"),
            str(FINDINGS / "gemini-flash-lite" / "provenance" / "provenance_summary.json"),
            str(FINDINGS / "glm-5" / "participation" / "participation_summary.json"),
            str(FINDINGS / "glm-5" / "provenance" / "provenance_summary.json"),
            str(FINDINGS / "kimi-k2.5" / "participation" / "participation_summary.json"),
            str(FINDINGS / "kimi-k2.5" / "provenance" / "provenance_summary.json"),
        ],
    )

    # Slide 14
    slide = prs.slides.add_slide(blank)
    add_title(slide, "The system narrows into a shared house style.", 14)
    add_card(slide, 0.8, 1.2, 8.35, 1.2, border=lighten(BLUE, 0.5), fill=lighten(BLUE, 0.95))
    add_textbox(
        slide,
        1.05,
        1.52,
        7.8,
        0.5,
        "The feed does not open into richer social behavior just because more agents are present. It repeatedly compresses into a locally shared procedural voice.",
        font_size=19,
        bold=True,
        color=GRAY_900,
        font_name=FONT_BOLD,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    boxes = [
        ("Topic changes survive", "Different seed worlds still push the discourse toward different vocabularies and topic anchors."),
        ("Personality does not disappear completely", "But it survives mostly as residue inside a much stronger shared format."),
        ("The social failure mode is predictable", "The network is more likely to settle into a template than to self-diversify."),
    ]
    xs = [0.78, 3.47, 6.16]
    for (title, detail), x in zip(boxes, xs):
        add_card(slide, x, 3.1, 2.38, 1.3, border=GRAY_200, fill=WHITE)
        add_textbox(slide, x + 0.08, 3.28, 2.08, 0.26, title, font_size=13, bold=True, color=GRAY_900)
        add_textbox(slide, x + 0.08, 3.68, 2.12, 0.46, detail, font_size=10, color=GRAY_700)
    register(
        metas,
        14,
        "The system narrows into a shared house style.",
        "State the broader interpretation plainly once the deck has shown the evidence.",
        "One large thesis card plus three supporting interpretation cards.",
        [
            "This is the synthesis slide, not a new-evidence slide.",
            "Use the phrase 'shared house style' verbally to make the result intuitive.",
        ],
        [],
    )

    # Slide 15
    slide = prs.slides.add_slide(blank)
    add_title(slide, "More agents can make synthetic publics more predictable, not more pluralistic.", 15)
    implication_cards = [
        ("Scale is not a diversity fix", "A bigger network can reinforce a dominant format instead of generating richer social variety.", BLUE),
        ("Synthetic publics can look larger than they are", "More accounts do not necessarily mean more viewpoints if the feed contracts into one shared style.", ORANGE),
        ("Interventions should target attractors", "Breaking collapse likely requires changing the social dynamics, not just adding more agents.", GREEN),
    ]
    for idx, (title, detail, accent) in enumerate(implication_cards):
        y = 1.18 + idx * 1.16
        add_card(slide, 0.72, y, 8.55, 0.9, border=lighten(accent, 0.45), fill=lighten(accent, 0.94))
        add_textbox(slide, 0.95, y + 0.1, 2.5, 0.2, title, font_size=14, bold=True, color=accent)
        add_textbox(slide, 3.08, y + 0.11, 5.85, 0.38, detail, font_size=12, color=GRAY_700)
    register(
        metas,
        15,
        "More agents can make synthetic publics more predictable, not more pluralistic.",
        "Translate the result from one experiment into why it matters for agentic social media more broadly.",
        "Three horizontal implication cards with short explanatory text.",
        [
            "Keep this tied to the evidence shown, not speculative grand theory.",
            "The main rhetorical move is: scale alone is not the road to rich emergent society.",
        ],
        [],
    )

    # Slide 16
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Limits now, useful next questions later.", 16)
    add_card(slide, 0.72, 1.2, 3.95, 2.8, border=GRAY_200, fill=WHITE)
    add_textbox(slide, 0.95, 1.36, 3.0, 0.22, "Limits", font_size=16, bold=True, color=GRAY_900)
    limit_lines = [
        "Controlled platform, not open internet behavior.",
        "One-hour windows, not long-horizon communities.",
        "GLM-5 and Kimi-K2.5 are n10-only here.",
    ]
    y = 1.78
    for line in limit_lines:
        add_textbox(slide, 0.95, y, 3.2, 0.28, f"• {line}", font_size=12, color=GRAY_700)
        y += 0.42
    add_card(slide, 5.05, 1.2, 4.15, 2.8, border=GRAY_200, fill=WHITE)
    add_textbox(slide, 5.28, 1.36, 3.2, 0.22, "Next questions", font_size=16, bold=True, color=GRAY_900)
    next_lines = [
        "Which ranking or incentive changes disrupt attractor formation?",
        "Does stronger heterogeneity delay or break collapse?",
        "What happens over longer social timescales?",
    ]
    y = 1.78
    for line in next_lines:
        add_textbox(slide, 5.28, y, 3.45, 0.28, f"• {line}", font_size=12, color=GRAY_700)
        y += 0.42
    add_card(slide, 1.35, 4.38, 7.3, 0.38, border=GRAY_200, fill=GRAY_100)
    add_textbox(
        slide,
        1.52,
        4.46,
        6.95,
        0.16,
        "The next question is not whether collapse exists. It is how to prevent it, slow it, or redirect it.",
        font_size=11,
        color=GRAY_700,
        align=PP_ALIGN.CENTER,
        margin=0.0,
    )
    register(
        metas,
        16,
        "Limits now, useful next questions later.",
        "Close the main deck with credibility and a clear research agenda.",
        "Two-column limits and next-questions layout plus one closing strip.",
        [
            "This is a clean research ending, not a dramatic ending.",
            "It should sound like the next experiment is now obvious and worth doing.",
        ],
        [],
    )

    # Slide 17 appendix divider
    slide = prs.slides.add_slide(blank)
    add_cover_title(slide, "Appendix", "Backup figures and extra examples", 17)
    register(
        metas,
        17,
        "Appendix",
        "Start the backup section for Q&A without cluttering the main narrative.",
        "Simple divider slide.",
        [],
        [],
    )

    # Slide 18
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Appendix: GPT-5 phrase adoption by scale", 18)
    add_picture_fit(slide, FINDINGS / "gpt-5" / "diffusion" / "scale_comparison.png", 0.45, 0.92, 8.95, 4.3)
    register(
        metas,
        18,
        "Appendix: GPT-5 phrase adoption by scale",
        "Keep the original multi-condition scale figure available for Q&A.",
        "Existing report figure placed nearly full-slide.",
        [],
        [str(FINDINGS / "gpt-5" / "diffusion" / "scale_comparison.png")],
    )

    # Slide 19
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Appendix: GPT-5 diversity collapse grid", 19)
    add_picture_fit(slide, FINDINGS / "gpt-5" / "diversity" / "diversity_grid.png", 0.58, 0.88, 8.8, 4.42)
    register(
        metas,
        19,
        "Appendix: GPT-5 diversity collapse grid",
        "Keep the full report-style diversity view available for detail questions.",
        "Existing report figure placed nearly full-slide.",
        [],
        [str(FINDINGS / "gpt-5" / "diversity" / "diversity_grid.png")],
    )

    # Slide 20
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Appendix: GPT-5 wave of conformity", 20)
    add_picture_fit(slide, FINDINGS / "gpt-5" / "participation" / "agent_usage_grid.png", 0.42, 0.88, 9.0, 4.42)
    register(
        metas,
        20,
        "Appendix: GPT-5 wave of conformity",
        "Keep the dense adoption-grid backup figure available for detailed discussion of within-run spread.",
        "Existing report figure placed nearly full-slide.",
        [],
        [str(FINDINGS / "gpt-5" / "participation" / "agent_usage_grid.png")],
    )

    # Slide 21
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Appendix: GPT-5 phrase DNA grid", 21)
    add_picture_fit(slide, FINDINGS / "gpt-5" / "provenance" / "phrase_dna_grid.png", 0.55, 0.98, 8.82, 4.15)
    register(
        metas,
        21,
        "Appendix: GPT-5 phrase DNA grid",
        "Keep the full run-by-run phrase-family backup figure available for provenance questions.",
        "Existing report figure placed nearly full-slide.",
        [],
        [str(FINDINGS / "gpt-5" / "provenance" / "phrase_dna_grid.png")],
    )

    # Slide 22
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Appendix: Example post families", 22)
    add_picture_fit(slide, FINDINGS / "gpt-5" / "provenance" / "heartbeat_ping_posts.png", 0.48, 1.08, 4.38, 3.75)
    add_picture_fit(slide, FINDINGS / "gpt-5" / "provenance" / "receipt_options_owner_posts.png", 5.0, 1.08, 4.38, 3.75)
    register(
        metas,
        22,
        "Appendix: Example post families",
        "Keep the qualitative backup examples available when people ask what the attractors actually look like.",
        "Two large existing post-family example images, side by side.",
        [],
        [
            str(FINDINGS / "gpt-5" / "provenance" / "heartbeat_ping_posts.png"),
            str(FINDINGS / "gpt-5" / "provenance" / "receipt_options_owner_posts.png"),
        ],
    )

    # Slide 23
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Appendix: GPT-5 topic concentration by phrase family", 23)
    add_picture_fit(slide, semantic_chart, 0.5, 0.98, 8.95, 4.15)
    register(
        metas,
        23,
        "Appendix: GPT-5 topic concentration by phrase family",
        "Keep the reduced semantic-concentration comparison available for questions about topic narrowing versus phrase-family narrowing.",
        "Rebuilt horizontal bar comparison figure covering representative AGI, tech-humor, and conspiracy runs.",
        [],
        [
            str(FINDINGS / "gpt-5" / "embedding_bridge" / "topic_anchor_run_summary.csv"),
            str(ASSET_DIR / "semantic_concentration.png"),
        ],
    )

    prs.save(PPTX_OUT)
    return metas


def write_outline(metas: list[SlideMeta]) -> None:
    lines = [
        "# Entropy Collapse Scaling Presentation Outline",
        "",
        f"Deck: `{PPTX_OUT.relative_to(ROOT)}`",
        "",
        "This outline mirrors the generated presentation and provides speaker-note bullets for each slide.",
        "",
    ]
    for meta in metas:
        lines.extend(
            [
                f"## Slide {meta.number}: {meta.title}",
                "",
                f"**Key message:** {meta.key_message}",
                "",
                f"**Visual spec:** {meta.visual_spec}",
                "",
                "**Speaker notes:**",
            ]
        )
        for note in meta.notes:
            lines.append(f"- {note}")
        if not meta.notes:
            lines.append("- No extra notes.")
        lines.append("")
    OUTLINE_OUT.write_text("\n".join(lines))


def write_manifest(metas: list[SlideMeta]) -> None:
    lines = [
        "# Entropy Collapse Scaling Presentation Manifest",
        "",
        f"Deck: `{PPTX_OUT.relative_to(ROOT)}`",
        "",
        "Each slide below lists the files or external sources that support its visible claims.",
        "",
    ]
    for meta in metas:
        lines.extend(
            [
                f"## Slide {meta.number}: {meta.title}",
                "",
                f"**Claim:** {meta.key_message}",
                "",
                "**Sources:**",
            ]
        )
        for source in meta.sources:
            lines.append(f"- `{source}`" if "://" not in source else f"- {source}")
        if not meta.sources:
            lines.append("- No explicit source file beyond slide-level synthesis and prior slides.")
        lines.append("")
    MANIFEST_OUT.write_text("\n".join(lines))


def main() -> None:
    metas = build_deck()
    write_outline(metas)
    write_manifest(metas)
    print(f"Wrote {PPTX_OUT}")
    print(f"Wrote {OUTLINE_OUT}")
    print(f"Wrote {MANIFEST_OUT}")


if __name__ == "__main__":
    main()
